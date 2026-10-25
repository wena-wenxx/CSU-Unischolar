"""
Reads a scholarship memo / announcement and SUGGESTS its eligibility rules
(ASSISTIVE ONLY). OAS staff see every suggestion with the sentence it came
from, tick the ones that are right, and press Save themselves. Nothing is
saved or approved here.

Two ways to read:
  1. A local large language model through Ollama (free, runs on the OAS
     computer, nothing leaves it). It uses the installed model
     (qwen2.5:1.5b or llama3.2:3b first) unless OLLAMA_MODEL is set.
       Install Ollama (ollama.com), then:  ollama pull llama3.2:3b
       Settings (.env of the AI service): OLLAMA_URL, OLLAMA_MODEL
     The model must quote the sentence for every value. Values whose quote
     is not found in the memo are marked "check": a language model can be
     wrong, so it is never trusted blindly.
  2. Without Ollama: simple patterns (GWA 2.50, income 300,000, "first year",
     document names, fields of study). Fewer details, but always works.
"""
import json
import os
import re
import urllib.request

from fit_matcher import fields_of

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")
# Empty = use whichever model is installed (qwen2.5:1.5b or llama3.2:3b first).
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "").strip()
PREFERRED_MODELS = ("qwen2.5:1.5b", "llama3.2:3b", "qwen2.5:3b", "llama3.2:1b")


def ollama_model(timeout=3):
    """The model to use: OLLAMA_MODEL if set, else the best one Ollama has installed."""
    if OLLAMA_MODEL:
        return OLLAMA_MODEL
    with urllib.request.urlopen(OLLAMA_URL.rstrip("/") + "/api/tags", timeout=timeout) as response:
        names = [m.get("name", "") for m in json.loads(response.read().decode()).get("models", [])]
    if not names:
        raise RuntimeError("Ollama is running but has no model. Run: ollama pull qwen2.5:1.5b")
    for wanted in PREFERRED_MODELS:
        if wanted in names:
            return wanted
    return names[0]
MAX_CHARS = 12000

DOCUMENT_PATTERNS = {
    "Certificate of Registration (COR)": [r"certificate of registration", r"\bcor\b", r"registration form"],
    "Certificate of Grades": [r"certificate of grades", r"\bgrades\b", r"transcript", r"report card", r"form 138"],
    "Certificate of Indigency": [r"indigency"],
    "Birth Certificate": [r"birth certificate"],
    "Valid ID": [r"valid id", r"school id", r"government[- ]issued id"],
    "Barangay Clearance": [r"barangay clearance"],
    "Certificate of Good Moral Character": [r"good moral"],
    "Parents' Income Tax Return": [r"income tax return", r"\bitr\b", r"bir form", r"tax exemption"],
    "Recommendation Letter": [r"recommendation letter", r"letter of recommendation"],
}
YEAR_PATTERNS = {
    "1st Year": [r"\bfirst[- ]year", r"\b1st[- ]year", r"incoming freshm", r"\bfreshm"],
    "2nd Year": [r"\bsecond[- ]year", r"\b2nd[- ]year"],
    "3rd Year": [r"\bthird[- ]year", r"\b3rd[- ]year"],
    "4th Year": [r"\bfourth[- ]year", r"\b4th[- ]year"],
}
MONTHS = "january february march april may june july august september october november december".split()


# ---------------------------------------------------------------- text helpers
def _norm(text):
    return re.sub(r"\s+", " ", (text or "")).strip().lower()


def _sentences(text):
    # Memos wrap lines in the middle of sentences: join them first, then split at . ; ! ?
    flat = re.sub(r"\s+", " ", text or "")
    return [s.strip() for s in re.split(r"(?<=[.;!?])\s+", flat) if s.strip()]


def _sentence_with(text, pattern):
    for sentence in _sentences(text):
        if re.search(pattern, sentence, re.I):
            return sentence[:300]
    return None


def _quote_found(quote, text):
    q = _norm(quote)
    return bool(q) and len(q) >= 6 and q in _norm(text)


def read_file_text(path, filename=""):
    """Plain text of a memo: PDF (text or OCR), Word .docx, .txt, or a picture (OCR)."""
    name = (filename or path).lower()
    if name.endswith(".txt"):
        with open(path, encoding="utf-8", errors="replace") as handle:
            return handle.read()
    if name.endswith(".docx"):
        import docx

        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    from ocr_engine import extract_text   # PDF and pictures (PaddleOCR for scans)

    result = extract_text(path, filename)
    if result.get("error"):
        raise ValueError(result["error"])
    return result.get("text", "")


# ---------------------------------------------------------------- 1. patterns
def pattern_suggest(text, options):
    s = {}
    low = re.sub(r"\s+", " ", text or "")   # one line: wrapped sentences still match

    m = re.search(r"(?:gwa|general weighted average|weighted average|grade)[^.\n]{0,80}?\b([1-3]\.\d{1,2})\b", low, re.I)
    if m:
        s["max_gwa"] = {"value": float(m.group(1)), "quote": _sentence_with(text, re.escape(m.group(1)))}

    m = re.search(r"income[^.\n]{0,100}?(?:php|₱|\bp)?\s?(\d{1,3}(?:,\d{3})+|\d{5,7})", low, re.I)
    if m:
        s["max_family_income"] = {"value": int(m.group(1).replace(",", "")), "quote": _sentence_with(text, re.escape(m.group(1)))}

    years = [label for label, pats in YEAR_PATTERNS.items() if any(re.search(p, low, re.I) for p in pats)]
    years = [y for y in years if y in options.get("year_levels", [y])]
    if years:
        first = next(p for y in years for p in YEAR_PATTERNS[y] if re.search(p, low, re.I))
        s["year_levels"] = {"value": years, "quote": _sentence_with(text, first)}

    memo_fields = fields_of(text)
    colleges = [c for c in options.get("colleges", []) if re.search(re.escape(c), low, re.I)
                or (memo_fields and set(fields_of(c)) & set(memo_fields))]
    if colleges:
        s["colleges"] = {"value": colleges, "quote": None}

    m = re.search(r"(no other|not (?:be )?(?:a )?(?:recipient|grantee|beneficiary|enjoying|availing)[^.\n]{0,40}(?:other|another)|only one scholarship)[^.\n]*", low, re.I)
    if m:
        s["one_scholarship_only"] = {"value": True, "quote": _sentence_with(text, re.escape(m.group(1)[:25]))}

    docs = [name for name, pats in DOCUMENT_PATTERNS.items() if any(re.search(p, low, re.I) for p in pats)]
    docs = [d for d in docs if d in options.get("documents", [d])]
    if docs:
        s["required_documents"] = {"value": docs, "quote": None}

    if memo_fields:
        s["fields_of_study"] = {"value": memo_fields, "quote": None}

    m = re.search(r"(?:deadline|on or before|until|not later than)[^.\n]{0,40}?\b(" + "|".join(MONTHS) + r")\s+(\d{1,2}),?\s+(\d{4})", low, re.I)
    if m:
        month = MONTHS.index(m.group(1).lower()) + 1
        s["deadline"] = {"value": f"{m.group(3)}-{month:02d}-{int(m.group(2)):02d}", "quote": _sentence_with(text, m.group(1) + r"\s+" + m.group(2))}

    m = re.search(r"(?:php|₱)\s?(\d{1,3}(?:,\d{3})+)(?:\.00)?\s*(?:per|a|/|each)\s*(?:semester|sem)\b", low, re.I)
    if m:
        s["amount"] = {"value": int(m.group(1).replace(",", "")), "quote": _sentence_with(text, re.escape(m.group(1)))}

    return s


# ---------------------------------------------------------------- 2. local LLM
PROMPT = """You read a Philippine scholarship memo and extract its eligibility rules.
Answer ONLY with JSON in exactly this shape. Use null when the memo does not say it. Do not guess.
For every value, "quote" must be the exact sentence from the memo that says it.
{{
 "max_gwa": {{"value": number on the 1.00 (best) to 5.00 scale or null, "quote": string or null}},
 "max_family_income": {{"value": yearly family income in pesos (number) or null, "quote": string or null}},
 "year_levels": {{"value": list from {years}, "quote": string or null}},
 "colleges": {{"value": list from {colleges}, "quote": string or null}},
 "one_scholarship_only": {{"value": true if the student must not hold another scholarship, else null, "quote": string or null}},
 "required_documents": {{"value": list from {documents}, "quote": string or null}},
 "fields_of_study": {{"value": list of short course areas the scholarship is for, "quote": string or null}},
 "deadline": {{"value": "YYYY-MM-DD" or null, "quote": string or null}},
 "amount": {{"value": pesos per semester (number) or null, "quote": string or null}}
}}

MEMO:
\"\"\"{memo}\"\"\""""


def _ollama(text, options, model, timeout=180):
    prompt = PROMPT.format(
        years=json.dumps(options.get("year_levels", [])),
        colleges=json.dumps(options.get("colleges", [])),
        documents=json.dumps(options.get("documents", [])),
        memo=text[:MAX_CHARS],
    )
    body = json.dumps({"model": model, "prompt": prompt, "stream": False, "format": "json",
                       "options": {"temperature": 0}}).encode()
    request = urllib.request.Request(OLLAMA_URL.rstrip("/") + "/api/generate", data=body,
                                     headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        answer = json.loads(response.read().decode())
    return json.loads(answer.get("response") or "{}")


# ---------------------------------------------------------------- cleaning
def _clean(raw, text, options, source):
    """Keep only valid values; list values must be from the allowed choices; check every quote."""
    out = {}
    allowed = {"year_levels": options.get("year_levels"), "colleges": options.get("colleges"),
               "required_documents": options.get("documents")}
    for key, item in (raw or {}).items():
        if not isinstance(item, dict):
            item = {"value": item, "quote": None}
        value, quote = item.get("value"), item.get("quote")
        if value in (None, "", [], False):
            continue
        try:
            if key == "max_gwa":
                value = round(float(value), 2)
                if not 1.0 <= value <= 5.0:
                    continue
            elif key in ("max_family_income", "amount"):
                value = int(float(str(value).replace(",", "")))
                if value <= 0:
                    continue
            elif key in allowed:
                value = [v for v in (value if isinstance(value, list) else [value]) if not allowed[key] or v in allowed[key]]
                if not value:
                    continue
            elif key == "fields_of_study":
                value = [str(v)[:60] for v in (value if isinstance(value, list) else [value])][:6]
            elif key == "one_scholarship_only":
                value = bool(value)
            elif key == "deadline":
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value)):
                    continue
            else:
                continue
        except (TypeError, ValueError):
            continue
        found = _quote_found(quote, text) if quote else None
        out[key] = {"value": value, "quote": quote if found else None, "by": source,
                    "check": source == "llm" and not found}
    return out


def suggest(text, options, use_llm=True):
    """Returns {method, suggestions, warnings}."""
    text = (text or "").strip()
    if len(text) < 20:
        raise ValueError("The memo text is too short to read.")
    warnings = []
    patterns = _clean(pattern_suggest(text, options), text, options, "pattern")

    if use_llm:
        model = OLLAMA_MODEL or "Ollama"
        try:
            model = ollama_model()
            llm = _clean(_ollama(text, options, model), text, options, "llm")
            # The model's answers first; the patterns fill what it left out.
            merged = {**patterns, **llm}
            return {"method": f"ollama:{model}", "suggestions": merged, "warnings": warnings}
        except Exception as e:
            reason = str(e) if isinstance(e, RuntimeError) else type(e).__name__
            warnings.append(f"The local AI model ({model}) is not running or failed ({reason}), "
                            "so simple pattern reading was used. Open the Ollama app, then try again.")

    return {"method": "patterns", "suggestions": patterns, "warnings": warnings}
