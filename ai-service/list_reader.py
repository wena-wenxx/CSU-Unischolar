"""
External grantee lists (ASSISTIVE ONLY).

Reads a list of scholarship grantees that an outside party (LGU, CHED,
private foundation) gave the OAS, and suggests which student in
ScholarGuide each line belongs to. OAS staff confirm or correct every
suggestion before anything is saved.

    read_table(path, filename)    -> rows of cells, whatever the file type
        .csv         Python's csv module
        .xlsx        openpyxl (every sheet)
        .docx        python-docx (tables; otherwise one row per paragraph)
        .pdf         PyMuPDF words, put back into rows and columns by their
                     position on the page; a scanned page (no text) goes
                     through PaddleOCR and is put back the same way
        .png / .jpg  PaddleOCR
    parse_records(rows)           -> one record per grantee: name + Student ID
    match_records(records, students, threshold)
                                  -> matched / check / unmatched / error,
                                     with the best candidates for staff

Matching rules (thresholds are 0-100 RapidFuzz scores):
    1. Student ID found in ScholarGuide -> that student, if the name on the
       line is similar enough (>= ID_NAME_MIN) or there is no name.
       ID found but the name looks different -> "check" (staff decide).
    2. No usable Student ID -> fuzzy name match (token_sort_ratio, so
       "DELA CRUZ, JUAN M." and "Juan M. Dela Cruz" are the same).
       best >= threshold (default 80) and clearly ahead of the 2nd best
       -> "matched"; two students almost equally close -> "check".
       If the line HAD a Student ID that is not in ScholarGuide, a name
       match is only a "check", never "matched".
    3. Nothing close enough -> "unmatched" (with suggestions >= SUGGEST_MIN).
    4. No name and no Student ID could be read -> "error".
"""
import csv
import os
import re
import zipfile
import statistics
import tempfile
import unicodedata

from rapidfuzz import fuzz

ID_NAME_MIN = 60      # Student ID found: the name must be at least this similar
SUGGEST_MIN = 60      # candidates shown to staff for unmatched / check rows
AMBIGUOUS_GAP = 3     # best and 2nd-best names closer than this -> "check"
MAX_ROWS = 5000

LIST_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".png", ".jpg", ".jpeg"}

# CSU-style IDs: 221-00462, 2024-10123 (spaces around the dash are tolerated).
ID_RE = re.compile(r"(?<![\d-])(\d{2,4})\s?-\s?(\d{4,6})(?![\d-])")

HEADER_PATTERNS = {
    "student_id": re.compile(r"^(student\s*(id|no\.?|number|#)|stud(ent)?\.?\s*no\.?|id\s*(no\.?|number)|school\s*id|id)$", re.I),
    "name": re.compile(r"^(name|full\s*name|complete\s*name|student'?s?\s*name|name\s*of\s*(the\s*)?(student|grantee|scholar|awardee)s?|grantee|scholar|awardee)$", re.I),
    "last": re.compile(r"^(last\s*name|surname|family\s*name|apelyido)$", re.I),
    "first": re.compile(r"^(first\s*name|given\s*name|pangalan)$", re.I),
    "middle": re.compile(r"^(middle\s*(name|initial)?|m\.?\s*i\.?)$", re.I),
}

# Cells that are clearly not a person's name (when there is no header row).
NOT_A_NAME = re.compile(
    r"\b(bachelor|bs[a-z]{0,4}|ab|college|university|department|year|sem(ester)?|course|program|"
    r"scholarship|grantee|list|total|page|republic|province|city|municipal|office|approved|amount|php|"
    r"no\.|number|date|signature|remarks|status|information|technology|science|sciences|education|"
    r"engineering|nursing|agriculture|forestry|accountancy|management|administration|arts|major)\b", re.I)


# --------------------------------------------------------------------- reading

def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "strftime"):          # Excel date -> 2026-10-22 (time only when there is one)
        value = value.strftime("%Y-%m-%d %H:%M" if getattr(value, "hour", 0) or getattr(value, "minute", 0) else "%Y-%m-%d")
    return re.sub(r"\s+", " ", str(value)).strip()


def _read_csv(path):
    with open(path, "rb") as fh:
        raw = fh.read()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[_clean(c) for c in row] for row in csv.reader(text.splitlines(), dialect)]


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:          # library not installed: use the built-in reader
        return _read_xlsx_builtin(path)

    try:
        # read_only=False: some files (Google Sheets, older tools) give wrong sizes in read-only mode.
        book = load_workbook(path, data_only=True)
    except zipfile.BadZipFile:
        raise ValueError("This is not a real .xlsx file (maybe an old .xls that was renamed). "
                         "Open it in Excel and use Save As .xlsx.")
    rows = []
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                rows.append([_clean(c) for c in row])
    finally:
        book.close()
    return rows


# Built-in readers (Python standard library only): a .xlsx / .docx file is a zip
# folder of XML files. Used when openpyxl / python-docx are not installed.
_X = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _open_zip(path, kind):
    try:
        return zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise ValueError(f"This is not a real .{kind} file (maybe an old file that was renamed). "
                         f"Open it and use Save As .{kind}.")


def _column(ref):
    letters = re.match(r"[A-Z]+", ref or "")
    if not letters:
        return None
    n = 0
    for ch in letters.group(0):
        n = n * 26 + ord(ch) - 64
    return n - 1


def _read_xlsx_builtin(path):
    import datetime
    from xml.etree import ElementTree as ET

    with _open_zip(path, "xlsx") as z:
        names = set(z.namelist())
        shared = []
        if "xl/sharedStrings.xml" in names:
            for si in ET.fromstring(z.read("xl/sharedStrings.xml")).iter(_X + "si"):
                shared.append("".join(t.text or "" for t in si.iter(_X + "t")))

        date_styles = set()
        if "xl/styles.xml" in names:
            styles = ET.fromstring(z.read("xl/styles.xml"))
            custom = set()
            for fmt in styles.iter(_X + "numFmt"):
                code = re.sub(r'"[^"]*"|\[[^\]]*\]', "", (fmt.get("formatCode") or "").lower())
                if re.search(r"[dy]|m{3,}", code):
                    custom.add(int(fmt.get("numFmtId")))
            xfs = styles.find(_X + "cellXfs")
            for i, xf in enumerate(xfs if xfs is not None else []):
                fid = int(xf.get("numFmtId") or 0)
                if 14 <= fid <= 22 or 45 <= fid <= 47 or fid in custom:
                    date_styles.add(i)

        sheets = []
        if "xl/workbook.xml" in names and "xl/_rels/workbook.xml.rels" in names:
            rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
            for sheet in ET.fromstring(z.read("xl/workbook.xml")).iter(_X + "sheet"):
                target = (rels.get(sheet.get(_R + "id")) or "").lstrip("/")
                if target:
                    sheets.append(target if target.startswith("xl/") else "xl/" + target)
        if not sheets:
            sheets = sorted(n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))

        rows = []
        for sheet in sheets:
            if sheet not in names:
                continue
            for row in ET.fromstring(z.read(sheet)).iter(_X + "row"):
                cells = []
                for c in row.findall(_X + "c"):
                    kind, v = c.get("t"), c.find(_X + "v")
                    raw = v.text if v is not None and v.text else ""
                    if kind == "s":
                        value = shared[int(raw)] if raw.isdigit() and int(raw) < len(shared) else ""
                    elif kind == "inlineStr":
                        value = "".join(t.text or "" for t in c.iter(_X + "t"))
                    elif kind == "b":
                        value = "TRUE" if raw == "1" else "FALSE"
                    elif kind in ("str", "e"):
                        value = raw
                    else:
                        value = raw
                        try:
                            number = float(raw)
                            if int(c.get("s") or 0) in date_styles:
                                value = (datetime.date(1899, 12, 30) + datetime.timedelta(days=round(number))).isoformat()
                            elif number.is_integer():
                                value = str(int(number))
                            else:
                                value = repr(number)        # "32.200000000000003" -> "32.2"
                        except ValueError:
                            pass
                    col = _column(c.get("r"))
                    col = len(cells) if col is None else col
                    cells.extend([""] * (col - len(cells)))
                    if col < len(cells):
                        cells[col] = _clean(value)
                    else:
                        cells.append(_clean(value))
                rows.append(cells)
    return rows


def _read_docx_builtin(path):
    from xml.etree import ElementTree as ET

    with _open_zip(path, "docx") as z:
        try:
            root = ET.fromstring(z.read("word/document.xml"))
        except KeyError:
            raise ValueError("This Word file has no text part. Open it in Word and Save As .docx again.")

    def text_of(node):
        out = []
        for el in node.iter():
            if el.tag == _W + "t":
                out.append(el.text or "")
            elif el.tag in (_W + "tab", _W + "br"):
                out.append("\t")
        return "".join(out)

    rows = []
    for tr in root.iter(_W + "tr"):
        cells = []
        for tc in tr.findall(_W + "tc"):
            cells.append(_clean(" ".join(text_of(p) for p in tc.iter(_W + "p"))))
            span = tc.find(f"{_W}tcPr/{_W}gridSpan")
            cells.extend([""] * (int(span.get(_W + "val")) - 1 if span is not None else 0))
        rows.append(cells)
    if rows:
        return rows, "table"
    for p in root.iter(_W + "p"):
        text = text_of(p)
        if text.strip():
            rows.append(_split_line(text))
    return rows, "text"



def _split_line(text):
    """A text line -> cells (tabs, or two or more spaces, separate columns)."""
    parts = [p.strip() for p in re.split(r"\t|\s{2,}|\s\|\s", text) if p.strip()]
    return parts or [text.strip()]


def _read_docx(path):
    try:
        import docx
    except ImportError:          # library not installed: use the built-in reader
        return _read_docx_builtin(path)

    try:
        document = docx.Document(path)
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise ValueError("This is not a real .docx file (maybe an old .doc that was renamed). "
                         "Open it in Word and use Save As .docx.")
    rows = []
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:   # merged cells repeat: keep one copy
                text = _clean(cell.text)
                if not cells or cells[-1] != text:
                    cells.append(text)
            rows.append(cells)
    if rows:
        return rows, "table"
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            rows.append(_split_line(paragraph.text))
    return rows, "text"


def group_boxes(boxes, with_positions=False):
    """
    Words with their position -> rows of cells.
    boxes: [(x0, y0, x1, y1, text), ...]  (PDF words or OCR text boxes)
    Words whose vertical centres are close form one row; inside a row, a gap
    wider than about one letter height starts a new cell (a new column).
    with_positions=True also returns where each cell starts and ends
    ([[(x0, x1), ...], ...]), used to put cells under the right header.
    """
    boxes = [b for b in boxes if str(b[4]).strip()]
    if not boxes:
        return ([], []) if with_positions else []
    heights = [max(1.0, b[3] - b[1]) for b in boxes]
    h = statistics.median(heights)

    rows = []
    for box in sorted(boxes, key=lambda b: ((b[1] + b[3]) / 2, b[0])):
        centre = (box[1] + box[3]) / 2
        if rows and abs(centre - rows[-1]["centre"]) <= h * 0.5:
            rows[-1]["boxes"].append(box)
            n = len(rows[-1]["boxes"])
            rows[-1]["centre"] += (centre - rows[-1]["centre"]) / n
        else:
            rows.append({"centre": centre, "boxes": [box]})

    out, layout = [], []
    for row in rows:
        cells, spans, current, start, last_x1 = [], [], [], None, None
        for box in sorted(row["boxes"], key=lambda b: b[0]):
            if last_x1 is not None and box[0] - last_x1 > h * 1.0:
                cells.append(" ".join(current))
                spans.append((start, last_x1))
                current, start = [], None
            current.append(str(box[4]).strip())
            start = box[0] if start is None else start
            last_x1 = box[2]
        if current:
            cells.append(" ".join(current))
            spans.append((start, last_x1))
        out.append([_clean(c) for c in cells])
        layout.append(spans)
    return (out, layout) if with_positions else out


def ocr_boxes(image_path):
    """PaddleOCR text boxes [(x0, y0, x1, y1, text)] (2.x and 3.x result formats)."""
    import ocr_engine  # loaded only when needed (PaddleOCR is slow to start)

    ocr = ocr_engine.get_ocr()
    if ocr is None:
        raise RuntimeError(ocr_engine._ocr_error or "PaddleOCR is not installed")
    try:
        result = ocr.predict(image_path) if hasattr(ocr, "predict") else ocr.ocr(image_path, cls=True)
    except Exception:
        result = ocr.ocr(image_path)

    boxes = []
    for page in result or []:
        if page is None:
            continue
        texts = polys = None
        try:  # 3.x
            texts = list(page["rec_texts"])
            polys = page.get("rec_boxes") if hasattr(page, "get") else page["rec_boxes"]
        except Exception:
            try:
                texts = list(page.json["res"]["rec_texts"])
                polys = page.json["res"].get("rec_boxes")
            except Exception:
                texts = None
        if texts is not None and polys is not None:
            for text, box in zip(texts, list(polys)):
                b = [float(v) for v in list(box)[:4]]
                boxes.append((b[0], b[1], b[2], b[3], text))
            continue
        if isinstance(page, list):  # 2.x: [[points, (text, score)], ...]
            for line in page:
                try:
                    points, (text, _score) = line[0], line[1]
                    xs = [p[0] for p in points]
                    ys = [p[1] for p in points]
                    boxes.append((min(xs), min(ys), max(xs), max(ys), text))
                except Exception:
                    pass
    return boxes


def _read_pdf(path):
    try:
        import pymupdf as fitz
    except ImportError:  # older PyMuPDF
        import fitz

    rows, layout, used_ocr = [], [], False
    with fitz.open(path) as doc:
        for page in doc:
            words = page.get_text("words")  # (x0, y0, x1, y1, word, block, line, word_no)
            if len(words) >= 3:
                page_rows, page_layout = group_boxes([(w[0], w[1], w[2], w[3], w[4]) for w in words], True)
                rows.extend(page_rows)
                layout.extend(page_layout)
                continue
            # Scanned page: render it and read it with OCR.
            pix = page.get_pixmap(dpi=200)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                tmp_path = tmp.name
            try:
                pix.save(tmp_path)
                page_rows, page_layout = group_boxes(ocr_boxes(tmp_path), True)
                rows.extend(page_rows)
                layout.extend(page_layout)
                used_ocr = True
            finally:
                os.remove(tmp_path)
    return rows, layout, "ocr" if used_ocr else "pdf-text"


def read_table(path, filename=""):
    """Returns {"rows": [[cell, ...]], "method": str, "warnings": [str]}."""
    ext = os.path.splitext(filename or path)[1].lower()
    warnings = []
    layout = None   # cell positions (PDF and pictures only)

    if ext == ".csv":
        rows, method = _read_csv(path), "csv"
    elif ext == ".xlsx":
        rows, method = _read_xlsx(path), "excel"
    elif ext == ".docx":
        rows, method = _read_docx(path)
        method = "word-" + method
    elif ext == ".pdf":
        rows, layout, method = _read_pdf(path)
    elif ext in (".png", ".jpg", ".jpeg"):
        rows, layout = group_boxes(ocr_boxes(path), True)
        method = "ocr"
    else:
        raise ValueError("Unsupported file type. Use PDF, Word (.docx), Excel (.xlsx), CSV, PNG or JPG.")

    if method == "ocr":
        warnings.append("This file is a scan or a picture, so it was read with OCR. "
                        "Letters and columns can be misread: check every row before saving.")

    if layout is None:
        layout = [None] * len(rows)
    kept = [(r, l) for r, l in zip(rows, layout) if any(c for c in r)]
    if len(kept) > MAX_ROWS:
        warnings.append(f"Only the first {MAX_ROWS} rows were read.")
        kept = kept[:MAX_ROWS]
    return {"rows": [r for r, _ in kept], "layout": [l for _, l in kept], "method": method, "warnings": warnings}


# --------------------------------------------------------------------- parsing

def normalize_id(text):
    """First Student ID in the text, as 2024-10123. School years (2026-2027) are skipped."""
    for m in ID_RE.finditer(text or ""):
        a, b = m.group(1), m.group(2)
        if len(a) == 4 and len(b) == 4 and int(b) == int(a) + 1:
            continue  # "AY 2026-2027" is a school year, not a Student ID
        return f"{a}-{b}"
    return ""


def normalize_name(text):
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z\s]", " ", text).upper()
    return re.sub(r"\s+", " ", text).strip()


def _find_header(row):
    """Column positions if this row is a header row, else None."""
    found = {}
    for index, cell in enumerate(row):
        label = re.sub(r"[:\s]+$", "", cell.strip())
        for key, pattern in HEADER_PATTERNS.items():
            if key not in found and pattern.match(label):
                found[key] = index
                break
    has_name = "name" in found or "last" in found or "first" in found
    if has_name and (len(found) >= 2 or "name" in found):
        return found
    return None


def _looks_like_name(cell):
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", cell)
    if len(letters) < 4 or re.search(r"[\d:_@]", cell) or NOT_A_NAME.search(cell):
        return False
    words = [w for w in re.split(r"[\s,]+", cell) if w]
    return 2 <= len(words) <= 7 or ("," in cell and len(words) >= 2)


def _guess(row):
    """No header: Student ID = a cell with an ID; name = a cell that looks like a name."""
    student_id = ""
    for cell in row:
        student_id = normalize_id(cell)
        if student_id:
            break
    names = [c for c in row if _looks_like_name(c)]
    # Prefer a "LAST, FIRST M." cell; otherwise the first name-like cell
    # (names usually come before the course or the remarks).
    with_comma = [c for c in names if "," in c]
    name = (with_comma or names or [""])[0]
    if not name:  # one-cell lines like "2024-10123 Dela Cruz, Juan M."
        joined = " ".join(row)
        rest = ID_RE.sub(" ", joined)
        rest = re.sub(r"^\s*\d+[.)]?\s+", "", rest).strip(" -|")
        if _looks_like_name(rest):
            name = rest
    return name, student_id


def _align(row, spans, header_starts):
    """PDF / picture row: put each cell under the header column it starts in
    (an empty cell, e.g. a missing Student ID, would otherwise shift the rest)."""
    cells = [""] * len(header_starts)
    for text, (x0, _x1) in zip(row, spans):
        column = 0
        for index, start in enumerate(header_starts):
            if start <= x0 + 6:
                column = index
        cells[column] = f"{cells[column]} {text}".strip()
    return cells


def parse_records(rows, layout=None):
    """Rows of cells -> [{"row", "name", "student_id", "raw", "note"}]."""
    layout = layout or [None] * len(rows)
    # A header row (Name, Student ID, ...) anywhere in the first 30 rows:
    # everything above it is a title (agency name, school year, ...).
    header, start, header_starts = None, 0, None
    for index, row in enumerate(rows[:30]):
        found = _find_header(row)
        if found:
            header, start = found, index + 1
            if layout[index] and len(layout[index]) == len(row):
                header_starts = [x0 for x0, _x1 in layout[index]]
            break

    records = []
    for number, (row, spans) in enumerate(zip(rows, layout), start=1):
        if number <= start:
            continue
        if header_starts and spans and len(spans) == len(row):
            row = _align(row, spans, header_starts)

        raw = " | ".join(c for c in row if c)
        if header:
            if _find_header(row):  # the header repeated at the top of the next page
                continue

            def get(key):
                i = header.get(key)
                return row[i].strip() if i is not None and i < len(row) else ""

            if "name" in header:
                name = get("name")
            else:
                last, first, middle = get("last"), get("first"), get("middle")
                name = f"{last}, {first} {middle}".strip(" ,") if last else f"{first} {middle}".strip()
            student_id = normalize_id(get("student_id")) if "student_id" in header else normalize_id(raw)
            if not name and not student_id:
                # a blank, total or repeated-title line inside the table
                if raw and not re.fullmatch(r"[\d\s.|]*", raw):
                    records.append({"row": number, "name": "", "student_id": "", "raw": raw,
                                    "note": "No name and no Student ID could be read in this line."})
                continue
        else:
            name, student_id = _guess(row)
            if not name and not student_id:
                continue  # titles, page numbers, signatures, ... (not a grantee line)

        records.append({"row": number, "name": name, "student_id": student_id, "raw": raw, "note": ""})
    return records, header is not None


# -------------------------------------------------------------------- matching

def _variants(student):
    first = normalize_name(student.get("first_name"))
    middle = normalize_name(student.get("middle_name"))
    last = normalize_name(student.get("last_name"))
    out = {f"{first} {last}".strip()}
    if middle:
        out.add(f"{first} {middle} {last}")
        out.add(f"{first} {middle[0]} {last}")
    return [v for v in out if v]


def name_score(name, student):
    target = normalize_name(name)
    if not target:
        return 0
    return round(max(fuzz.token_sort_ratio(target, v) for v in _variants(student)))


def _display(student):
    return " ".join(p for p in [student.get("first_name"), student.get("middle_name"), student.get("last_name")] if p)


def _ranked(name, students, limit=3):
    scored = [(name_score(name, s), s) for s in students]
    scored.sort(key=lambda pair: (-pair[0], _display(pair[1])))
    return scored[:limit]


def _suggestion(score, student):
    return {"id": student.get("id"), "student_id": student.get("student_id"), "name": _display(student), "score": int(score)}


def match_records(records, students, threshold=80):
    by_id = {}
    for s in students:
        sid = normalize_id(s.get("student_id", "")) or str(s.get("student_id", "")).strip()
        if sid:
            by_id[sid.upper()] = s

    used = {}
    out = []
    for rec in records:
        result = dict(rec, result="unmatched", method=None, score=None, student=None, suggestions=[])
        name, sid = rec.get("name", ""), rec.get("student_id", "")

        if not name and not sid:
            result.update(result="error", note=rec.get("note") or "No name and no Student ID could be read in this line.")
            out.append(result)
            continue

        chosen = None
        if sid and sid.upper() in by_id:
            student = by_id[sid.upper()]
            score = name_score(name, student) if name else None
            result["suggestions"] = [_suggestion(score if score is not None else 100, student)]
            if score is None or score >= ID_NAME_MIN:
                chosen = student
                result.update(result="matched", method="student_id", score=100 if score is None else score)
            else:
                result.update(result="check", method="student_id", score=score,
                              note=f"The Student ID belongs to {_display(student)}, but the name on the list looks different.")
                result["student"] = _suggestion(score, student)
        elif name:
            ranked = _ranked(name, students)
            result["suggestions"] = [_suggestion(sc, st) for sc, st in ranked if sc >= SUGGEST_MIN]
            if ranked and ranked[0][0] >= threshold:
                best_score, best = ranked[0]
                close = len(ranked) > 1 and ranked[1][0] >= threshold and best_score - ranked[1][0] < AMBIGUOUS_GAP
                if close:
                    result.update(result="check", method="name", score=best_score,
                                  note="Two or more students have a very similar name. Pick the right one.")
                    result["student"] = _suggestion(best_score, best)
                elif sid:
                    result.update(result="check", method="name", score=best_score,
                                  note=f"Student ID {sid} is not in ScholarGuide; the name matches {_display(best)}.")
                    result["student"] = _suggestion(best_score, best)
                else:
                    chosen = best
                    result.update(result="matched", method="name", score=best_score)
            else:
                result.update(note=(f"Student ID {sid} is not in ScholarGuide and no" if sid else "No")
                              + f" student's name is {threshold}% similar.")
        else:
            result.update(note=f"Student ID {sid} is not in ScholarGuide, and there is no name to compare.")

        if chosen is not None:
            result["student"] = _suggestion(result["score"], chosen)
            key = chosen.get("id")
            if key in used:
                result.update(result="check", note=f"Same student as line {used[key]} (listed twice?).")
            else:
                used[key] = rec["row"]
        out.append(result)
    return out


def summarize(rows):
    summary = {"total": len(rows), "matched": 0, "check": 0, "unmatched": 0, "error": 0}
    for r in rows:
        summary[r["result"]] = summary.get(r["result"], 0) + 1
    return summary


def extract_and_match(path, filename, students, threshold=80):
    table = read_table(path, filename)
    records, has_header = parse_records(table["rows"], table.get("layout"))
    warnings = list(table["warnings"])
    if not has_header and records:
        warnings.append("No header row (Name, Student ID) was found, so the columns were guessed.")
    if not records:
        warnings.append("No grantee lines were found in this file.")
    rows = match_records(records, students, threshold)
    return {"method": table["method"], "warnings": warnings, "rows": rows, "summary": summarize(rows)}


def table_rows(path, filename):
    """
    Any list file -> plain rows of cells, for pages that let staff pick the
    columns themselves (Approved Lists, Enrollment). For PDFs and pictures,
    cells are put under the columns of the widest row near the top (the
    header), so an empty cell does not shift the rest of the row.
    """
    table = read_table(path, filename)
    rows, layout = table["rows"], table.get("layout") or []
    if rows and any(layout):
        top = [(len(r), i) for i, r in enumerate(rows[:30]) if layout[i]]
        if top:
            width, index = max(top, key=lambda pair: (pair[0], -pair[1]))
            starts = [x0 for x0, _x1 in layout[index]]
            if width >= 2:
                rows = [
                    _align(row, layout[i], starts) if i > index and layout[i] and len(layout[i]) == len(row) else row
                    for i, row in enumerate(rows)
                ]
    return {"rows": rows, "method": table["method"], "warnings": table["warnings"]}
